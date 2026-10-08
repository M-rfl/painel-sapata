"""Validation of a conversion: golden checks, workbook comparison and markdown report.

This is the ONLY module that touches openpyxl (imported lazily in read_workbook), so
the converter itself runs without it. See .planning/research/WORKBOOK-CRITERIA.md.
"""
from __future__ import annotations

import unicodedata
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal

from .analysis_ref import reference_analysis
from .gps import sample_key
from .model import ACCEL_PER_G, StructuralError

DEFAULT_FOLDED = {"Central – Dentro do Túnel": "Central – Gamboa"}
CLASS_LABELS = {"normal": 0, "leve": 1, "moderado": 2, "severo": 3, "critico": 4}
CLASS_NAMES = ["Normal", "Leve", "Moderado", "Severo", "Crítico"]


# ---------------------------------------------------------------- formatting

def pt_br_int(n) -> str:
    return f"{int(n):,}".replace(",", ".")


def pt_br_float(x, decimals: int = 2) -> str:
    s = f"{x:,.{decimals}f}"
    return s.replace(",", "\0").replace(".", ",").replace("\0", ".")


def _fmt(v) -> str:
    if v is None:
        return "—"
    if isinstance(v, bool):
        return str(v)
    if isinstance(v, int):
        return pt_br_int(v)
    if isinstance(v, float):
        return pt_br_int(v) if v == int(v) else pt_br_float(v, 1)
    if isinstance(v, dict):
        return ", ".join(f"{k}: {_fmt(x)}" for k, x in v.items())
    if isinstance(v, (list, tuple)):
        return "[" + ", ".join(_fmt(x) for x in v) + "]"
    return str(v)


# ---------------------------------------------------------------- golden checks

@dataclass
class Check:
    name: str
    expected: object
    actual: object
    ok: bool
    exact: bool
    note: str = ""


def _same(a, b) -> bool:
    if isinstance(a, (list, tuple)) and isinstance(b, (list, tuple)):
        return len(a) == len(b) and all(_same(x, y) for x, y in zip(a, b))
    if isinstance(a, dict) and isinstance(b, dict):
        return a.keys() == b.keys() and all(_same(a[k], b[k]) for k in a)
    if isinstance(a, (int, float)) and isinstance(b, (int, float)):
        return abs(a - b) < 1e-9
    return a == b


def golden_checks(ref: dict, counts: dict, golden: dict) -> list[Check]:
    ex = golden.get("exact", {})
    rf = golden.get("reference", {})
    actual_exact = {
        "rows": counts.get("rows"),
        "rows_per_file": counts.get("rows_per_file"),
        "rows_with_gps": counts.get("rows_with_gps"),
        "rows_without_gps": counts.get("rows_without_gps"),
        "medians_lsb": ref.get("medians_lsb"),
        "class_counts": ref.get("class_counts"),
        "eventos": ref.get("eventos"),
        "impactos": ref.get("impactos"),
        "impacts_by_class": ref.get("impacts_by_class"),
    }
    actual_ref = {
        "pontos": ref.get("pontos"),
        "points_by_class": ref.get("points_by_class"),
        "sentido_records": counts.get("sentido_records"),
        "sentido_eventos": counts.get("sentido_eventos"),
        "sentido_split_index": counts.get("sentido_split_index"),
        "trecho_most_events": counts.get("trecho_most_events"),
    }
    out: list[Check] = []
    for k, actual in actual_exact.items():
        if k in ex:
            out.append(Check(k, ex[k], actual, _same(ex[k], actual), True))
    for k, actual in actual_ref.items():
        if k in rf:
            out.append(Check(k, rf[k], actual, _same(rf[k], actual), False))
    return out


# ---------------------------------------------------------------- workbook

def _norm(text) -> str:
    s = unicodedata.normalize("NFKD", str(text if text is not None else ""))
    return "".join(c for c in s if not unicodedata.combining(c)).strip().lower()


def _col_matchers():
    def starts(p, *, not_in=()):
        return lambda h: h.startswith(p) and not any(x in h for x in not_in)

    return {
        "time": starts("data/hora"),
        "ax": starts("ax", not_in=("din.", "media")),
        "ay": starts("ay", not_in=("din.", "media")),
        "az": starts("az", not_in=("din.", "media")),
        "lat": starts("latitude"),
        "lon": starts("longitude"),
        "local": starts("local (chave"),
        "trecho": starts("trecho (padronizado"),
        "sentido": lambda h: h == "sentido",
        "cls": starts("classificacao"),
        "evento": starts("evento (1/0)"),
    }


_REQUIRED = ("ax", "ay", "az", "lat", "lon", "trecho", "cls", "evento")


def read_workbook(path) -> dict:
    """Stream sheet '02…' of the treated workbook into per-column lists."""
    import openpyxl  # lazy: only validate needs it

    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    try:
        sheets = [w for w in wb.worksheets if str(w.title).startswith("02")]
        if not sheets:
            raise StructuralError(f"{path}: no sheet whose name starts with '02' (found {wb.sheetnames})")
        ws = sheets[0]
        rows = ws.iter_rows(values_only=True)
        head_row = None
        buffered = []
        for r in rows:
            buffered.append(r)
            if any(_norm(c).startswith("classificacao") for c in r):
                head_row = r
                break
            if len(buffered) >= 10:
                break
        if head_row is None:
            raise StructuralError(f"{path}: header row with 'CLASSIFICAÇÃO' not found in first 10 rows")
        headers = [_norm(c) for c in head_row]
        cols: dict[str, int] = {}
        for key, fn in _col_matchers().items():
            for i, h in enumerate(headers):
                if h and fn(h):
                    cols[key] = i
                    break
        missing = [k for k in _REQUIRED if k not in cols]
        if missing:
            raise StructuralError(
                f"{path}: missing workbook columns {missing}; found headers: "
                f"{[str(c) for c in head_row if c is not None]}")
        out: dict[str, list] = {k: [] for k in cols}
        for r in rows:
            if r is None:
                continue
            if all(i >= len(r) or r[i] is None for i in (cols["cls"], cols["ax"])):
                continue
            for k, i in cols.items():
                out[k].append(r[i] if i < len(r) else None)
        out["n"] = len(out["cls"])
        return out
    finally:
        wb.close()


def _wb_time_second(v):
    if isinstance(v, datetime):
        return v.replace(microsecond=0, tzinfo=None)
    if isinstance(v, str):
        for f in ("%d/%m/%Y %H:%M:%S", "%Y-%m-%d %H:%M:%S", "%Y-%m-%dT%H:%M:%S"):
            try:
                return datetime.strptime(v.strip()[:19], f)
            except ValueError:
                continue
    return None


def _num(v):
    if v is None or v == "":
        return None
    try:
        return float(v)
    except (TypeError, ValueError):
        return None


def _wb_key(lat, lon):
    la, lo = _num(lat), _num(lon)
    if la is None or lo is None:
        return None
    return sample_key(Decimal(repr(la)), Decimal(repr(lo)))


def _wb_class(v):
    return CLASS_LABELS.get(_norm(v))


def _runs_of(names):
    out = []
    prev = object()
    for i, nm in enumerate(names):
        if nm != prev:
            out.append((i, nm))
            prev = nm
    return out


def workbook_comparison(run, assignment, decoded, ref, wb, *, folded=None) -> dict:
    folded = DEFAULT_FOLDED if folded is None else folded
    n_conv, n_wb = run.n, wb["n"]
    n = min(n_conv, n_wb)

    # row-level mismatches
    acc_mis = {"ax": 0, "ay": 0, "az": 0}
    for a in acc_mis:
        conv = decoded[a]
        for i in range(n):
            w = _num(wb[a][i])
            c = conv[i]
            if w is None or c is None:
                if w is not None or c is not None:
                    acc_mis[a] += 1
            elif abs(w - c / ACCEL_PER_G) > 1e-9:
                acc_mis[a] += 1
    gps_mis = 0
    for i in range(n):
        if _wb_key(wb["lat"][i], wb["lon"][i]) != decoded["loc_key"][i]:
            gps_mis += 1
    wb_cls = [_wb_class(v) for v in wb["cls"]]
    class_mis = sum(1 for i in range(n) if wb_cls[i] != ref["class_per_row"][i])
    time_mis = None
    if "time" in wb:
        checked = 0
        time_mis = 0
        for i in range(n):
            w = _wb_time_second(wb["time"][i])
            if w is None:
                continue
            checked += 1
            if w != run.timestamps[i].replace(microsecond=0, tzinfo=None):
                time_mis += 1
        if checked == 0:
            time_mis = None

    # workbook own analysis (same rules, workbook keys/trechos) for per-trecho numbers
    lsb = {a: [None if _num(v) is None else round(_num(v) * ACCEL_PER_G) for v in wb[a]]
           for a in ("ax", "ay", "az")}
    wb_keys = [_wb_key(la, lo) for la, lo in zip(wb["lat"], wb["lon"])]
    wb_names = [str(t) if t is not None else "" for t in wb["trecho"]]
    raw = reference_analysis(lsb["ax"], lsb["ay"], lsb["az"], wb_keys, wb_names)["per_trecho"]
    zero = {"records": 0, "eventos": 0, "impactos": 0, "pontos": 0}
    wb_pt: dict[str, dict] = {}
    detail = []
    for name, v in raw.items():
        tgt = folded.get(name, name)
        d = wb_pt.setdefault(tgt, dict(zero))
        for k in zero:
            d[k] += v[k]
    for src, dst in folded.items():
        if src in raw:
            detail.append({"from": src, "into": dst, "from_records": raw[src]["records"],
                           "into_records": raw.get(dst, zero)["records"],
                           "sum_records": wb_pt[dst]["records"],
                           "conv_records": ref["per_trecho"].get(dst, zero)["records"]})
    order = list(wb_pt)
    for name in ref["per_trecho"]:
        if name not in wb_pt:
            order.append(name)
    table = []
    for name in order:
        w = wb_pt.get(name, zero)
        c = ref["per_trecho"].get(name, zero)
        table.append({"trecho": name, "wb": dict(w), "conv": dict(c),
                      "delta": {k: c[k] - w[k] for k in zero}})

    # boundaries: converter trecho_run starts vs workbook run starts of the same name
    wb_runs = _runs_of([folded.get(x, x) for x in wb_names])
    seg_starts = [s for s, _ in assignment.sentido_runs] + [run.n]
    boundaries = []
    for k, (start, tid) in enumerate(assignment.trecho_runs):
        if k == 0:
            continue
        td = assignment.trechos[tid]
        seg_end = next((e for e in seg_starts if e > start), run.n)
        seg_start = max(s for s in seg_starts[:-1] if s <= start)
        cand = [s for s, nm in wb_runs if nm == td.name and seg_start <= s < seg_end]
        wb_idx = cand[0] if cand else None
        boundaries.append({"station": td.from_code, "trecho": td.name, "conv_index": start,
                           "wb_index": wb_idx,
                           "delta": None if wb_idx is None else start - wb_idx})

    return {
        "rows_wb": n_wb, "rows_conv": n_conv,
        "accel_mismatch": acc_mis, "gps_mismatch": gps_mis, "class_mismatch": class_mis,
        "time_mismatch": time_mis,
        "per_trecho": table, "folded": detail, "boundaries": boundaries,
        "wb_pontos": sum(v["pontos"] for v in wb_pt.values()),
    }


def workbook_checks(cmp: dict) -> list[Check]:
    am = cmp["accel_mismatch"]
    return [
        Check("workbook: linhas", cmp["rows_wb"], cmp["rows_conv"], cmp["rows_wb"] == cmp["rows_conv"], True),
        Check("workbook: AX/AY/AZ divergentes", 0, sum(am.values()), sum(am.values()) == 0, True,
              f"ax={am['ax']}, ay={am['ay']}, az={am['az']}"),
        Check("workbook: chaves GPS (6 casas) divergentes", 0, cmp["gps_mismatch"], cmp["gps_mismatch"] == 0, True),
        Check("workbook: classe por linha divergente", 0, cmp["class_mismatch"], cmp["class_mismatch"] == 0, True),
    ]


# ---------------------------------------------------------------- report

def _table(head, rows) -> list[str]:
    out = ["| " + " | ".join(head) + " |", "|" + "|".join("---" for _ in head) + "|"]
    for r in rows:
        out.append("| " + " | ".join(str(c) for c in r) + " |")
    return out


def render_report(ctx: dict) -> str:
    L: list[str] = []
    ref = ctx["ref"]
    a = ctx["assignment"]
    counts = ctx.get("counts", {})
    cmp = ctx.get("comparison")
    L.append(f"# Relatório de validação — {ctx['run_id']}")
    L.append("")
    L.append(f"- Gerado em: {ctx.get('generated_at', '')}")
    L.append(f"- Linha detectada: {ctx['line_id']}"
             + (" (AMBÍGUA)" if ctx.get("ambiguous") else "")
             + " — pontuações: "
             + ", ".join(f"linha {i}: {pt_br_float(s, 1)}" for i, s in ctx.get("line_scores", [])))
    L.append(f"- Raio da geocerca: {pt_br_float(ctx['radius_m'], 0)} m; raio nos terminais: "
             f"{pt_br_float(ctx['terminal_radius_m'], 0)} m")
    L.append("- Arquivos de origem:")
    for s in ctx.get("sources", []):
        L.append(f"  - `{s.name}` — {pt_br_int(s.rows)} linhas, sha256 `{s.sha256}`")
    for note in a.notes:
        L.append(f"- Nota: {note}")
    L.append("")

    L.append("## Verificação exata")
    L.append("")
    exact = [c for c in ctx["checks"] if c.exact]
    L.extend(_table(["Métrica", "Planilha/Golden", "Conversor", "Status"],
                    [(c.name, _fmt(c.expected), _fmt(c.actual) + (f" ({c.note})" if c.note else ""),
                      "OK" if c.ok else "FALHA") for c in exact]) if exact else ["Nenhum golden informado."])
    L.append("")

    L.append("## Pontos e trechos (lado a lado)")
    L.append("")
    refc = [c for c in ctx["checks"] if not c.exact]
    L.append(f"- Pontos de impacto (conversor): {pt_br_int(ref['pontos'])}; por classe: {_fmt(ref['points_by_class'])}")
    if cmp:
        L.append(f"- Pontos recalculados sobre as chaves da planilha: {pt_br_int(cmp['wb_pontos'])}")
    if refc:
        L.append("")
        L.extend(_table(["Referência", "Planilha", "Conversor", "Diferença"],
                        [(c.name, _fmt(c.expected), _fmt(c.actual), "igual" if c.ok else "difere")
                         for c in refc]))
    L.append("")
    if cmp:
        L.append("### Por trecho")
        L.append("")
        rows = []
        for r in cmp["per_trecho"]:
            w, c, d = r["wb"], r["conv"], r["delta"]
            rows.append((r["trecho"],
                         *(f"{pt_br_int(w[k])} / {pt_br_int(c[k])} ({d[k]:+d})"
                           for k in ("records", "eventos", "impactos", "pontos"))))
        L.extend(_table(["Trecho", "Registros (plan/conv)", "Eventos", "Impactos", "Pontos"], rows))
        L.append("")
        L.append("### Fronteiras entre trechos")
        L.append("")
        L.extend(_table(["Estação", "Trecho seguinte", "Índice planilha", "Índice conversor", "Delta (amostras)"],
                        [(b["station"], b["trecho"], _fmt(b["wb_index"]), pt_br_int(b["conv_index"]),
                          "—" if b["delta"] is None else f"{b['delta']:+d}") for b in cmp["boundaries"]]))
        L.append("")
        L.append("### Linha a linha")
        L.append("")
        am = cmp["accel_mismatch"]
        L.append(f"- Linhas planilha/conversor: {pt_br_int(cmp['rows_wb'])} / {pt_br_int(cmp['rows_conv'])}")
        L.append(f"- Aceleração divergente: AX {pt_br_int(am['ax'])}, AY {pt_br_int(am['ay'])}, AZ {pt_br_int(am['az'])}")
        L.append(f"- Chave GPS (6 casas) divergente: {pt_br_int(cmp['gps_mismatch'])}")
        L.append(f"- Classe divergente: {pt_br_int(cmp['class_mismatch'])}")
        tm = cmp["time_mismatch"]
        L.append("- Segundo do horário divergente: " + ("não verificado" if tm is None else pt_br_int(tm)))
        L.append("")
    L.append("### Explicações")
    L.append("")
    folded_map = ctx.get("folded") or DEFAULT_FOLDED
    for src, dst in folded_map.items():
        if cmp:
            d = next((x for x in cmp["folded"] if x["from"] == src), None)
        else:
            d = None
        if d:
            L.append(f"- Trecho de túnel dobrado: '{src}' ({pt_br_int(d['from_records'])}) + '{dst}' "
                     f"({pt_br_int(d['into_records'])}) = {pt_br_int(d['sum_records'])} registros na planilha, "
                     f"contra {pt_br_int(d['conv_records'])} no conversor (16+17 ≈ 16): o conversor não tem trecho próprio para o túnel.")
        else:
            L.append(f"- Trecho de túnel dobrado: '{src}' é somado a '{dst}' (16+17 ≈ 16); o conversor não tem trecho próprio para o túnel.")
    L.append("- Fronteiras: a planilha usa limites manuais; o conversor usa geocerca (raio por estação), então os índices de fronteira diferem em algumas amostras.")
    L.append("- Resolução de tempo: o conversor mantém carimbos sub-segundo; a planilha trabalha com resolução de 1 s, o que muda o agrupamento por segundo e os pontos.")
    L.append("- Estações inferidas (sem passagem dentro do raio) deslocam a fronteira do trecho vizinho.")
    L.append(f"- Linhas sem GPS: {pt_br_int(counts.get('rows_without_gps', 0))} no conversor; o rótulo 'Dentro do Túnel' "
             "da planilha cobre 882 linhas, das quais só 108 não têm coordenadas.")
    L.append("")

    L.append("## Sentidos")
    L.append("")
    L.append(f"- Registros por sentido: {_fmt(counts.get('sentido_records'))}; eventos: {_fmt(counts.get('sentido_eventos'))}")
    gold_split = next((c.expected for c in ctx["checks"] if c.name == "sentido_split_index"), None)
    L.append(f"- Índice da virada: {_fmt(counts.get('sentido_split_index'))}"
             + (f" (planilha: {_fmt(gold_split)})" if gold_split is not None else ""))
    for t in a.turnarounds:
        L.append(f"- Virada no índice {pt_br_int(t.index)}: método {t.method}, terminal {t.terminal_code}, "
                 f"maior intervalo {'—' if t.gap_s is None else pt_br_float(t.gap_s, 1) + ' s'}")
    L.append("")

    L.append("## Estações")
    L.append("")
    if a.passages:
        L.extend(_table(["Sentido", "Segmento", "Estação", "Status", "Distância mínima (m)"],
                        [(p.sentido + 1, p.segment + 1, p.code, p.status,
                          "—" if p.min_dist_m is None else pt_br_float(p.min_dist_m, 1)) for p in a.passages]))
        L.append("")
    L.append("**Estações nunca atingidas:** "
             + ("nenhuma" if not a.stations_not_hit else "; ".join(
                 f"sentido {s + 1} {c} (mín. {'—' if d is None else pt_br_float(d, 1)} m)"
                 for s, c, d in a.stations_not_hit)))
    inferred = [p for p in a.passages if p.status == "inferred"]
    L.append("")
    L.append("**Estações inferidas:** "
             + ("nenhuma" if not inferred else "; ".join(
                 f"sentido {p.sentido + 1} {p.code} (mín. {'—' if p.min_dist_m is None else pt_br_float(p.min_dist_m, 1)} m)"
                 for p in inferred)))
    L.append("")
    L.append("**Avisos dos arquivos de estações:** " + ("nenhum" if not ctx.get("station_warnings") else ""))
    for w in ctx.get("station_warnings", []):
        L.append(f"- [{w.kind}] {w.message}")
    L.append("")

    L.append("## Valores recuperados")
    L.append("")
    stats = ctx.get("recovery_stats") or {}
    if stats:
        L.extend(_table(["Estatística", "Valor"], [(k, pt_br_int(v)) for k, v in sorted(stats.items())]))
        L.append("")
    issues = ctx.get("issues", [])
    unrec = [i for i in issues if i.value is None]
    amb = [i for i in issues if i.reason == "ambiguous_resolved_by_rule"]
    L.append(f"**Valores não recuperados ({pt_br_int(len(unrec))}):**" + ("" if unrec else " nenhum"))
    for i in unrec:
        L.append(f"- `{i.file}` linha {pt_br_int(i.line)}, coluna `{i.column}`: `{i.raw}` ({i.reason})")
    L.append("")
    L.append(f"**Valores ambíguos resolvidos por regra ({pt_br_int(len(amb))}):**" + ("" if amb else " nenhum"))
    for i in amb:
        L.append(f"- `{i.file}` linha {pt_br_int(i.line)}, coluna `{i.column}`: `{i.raw}` -> {_fmt(i.value)}")
    L.append("")
    return "\n".join(L)
